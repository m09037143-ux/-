from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from repair_report.analytics.engine import build_report

FIXTURES = Path(__file__).parent / "fixtures"
MULTI_MONTH_FILE = FIXTURES / "WR_Consolidated_List_20260902_весь.xlsx"
JULY_ONLY_FILE = FIXTURES / "WR_Consolidated_List_20260902_июль.xlsx"
KOSOVOV_FILE = FIXTURES / "Final_Report_2026_июль_Косовов.xlsx"


@pytest.fixture(scope="session")
def july_report():
    """Report built from the multi-month file, July selected as current
    (June auto-resolves as previous). This is the primary fixture nearly
    every test validates against."""
    return build_report(str(MULTI_MONTH_FILE))


@pytest.fixture(scope="session")
def july_only_report():
    """Report built from the single-month (July-only) file -- exercises the
    'no previous period available' path."""
    return build_report(str(JULY_ONLY_FILE))


@pytest.fixture(scope="session")
def kosovov_sheets():
    """dict[sheet_name] -> list[tuple] (data rows only, header excluded).

    This workbook is a DEV-ONLY TEST FIXTURE -- see repair_report/config's
    module docs and docs/REVERSE_ENGINEERING.md §0. Nothing under
    repair_report/{ingest,analytics,ui} may read this file or one shaped
    like it; only this test suite does.
    """
    wb = openpyxl.load_workbook(str(KOSOVOV_FILE), data_only=True)
    result = {}
    for name in wb.sheetnames:
        ws = wb[name]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        result[name] = rows
    return result
